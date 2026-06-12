#!/usr/bin/env python3
"""Extract every CELNET.INSTRUMENT cell's resolved (underlier, product, terms, tenor,
notional) from a generated workbook → /tmp/wb_instruments.json. Resolves cell refs +
ranges via openpyxl so the REAL shapeSpecInstrument (scripts/_wbShape.ts) can validate
the terms. Run: /tmp/celnet-xlsx-venv/bin/python excel/scripts/_wbExtract.py <xlsx>"""
import json, re, sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/wb_check.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)


def cellval(ws, ref):
    v = ws[ref].value
    return "" if v is None else v


out = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith("=CELNET.INSTRUMENT(")):
                continue
            inner = v[v.index("(") + 1 : v.rindex(")")]
            args = []
            cur = ""
            q = False
            for ch in inner:
                if ch == '"':
                    q = not q
                if ch == "," and not q:
                    args.append(cur.strip())
                    cur = ""
                    continue
                cur += ch
            args.append(cur.strip())

            def lit(a, ws=ws):
                a = a.strip()
                if a.startswith('"'):
                    return a.strip('"')
                if re.match(r"^\$?[A-Z]+\$?\d+$", a):
                    return cellval(ws, a.replace("$", ""))
                return a

            def rng(a, ws=ws):
                a = a.replace("$", "")
                rows = []
                for r in ws[a]:
                    rr = ["" if x.value is None else x.value for x in r]
                    while rr and rr[-1] == "":
                        rr.pop()
                    if rr:
                        rows.append(rr)
                return rows

            underlier = lit(args[0])
            product = lit(args[1])
            terms = rng(args[2]) if ":" in args[2] else [[lit(args[2])]]
            tenor = lit(args[3]) if len(args) > 3 else None
            notional = lit(args[4]) if len(args) > 4 else None
            out.append({
                "sheet": ws.title, "cell": c.coordinate,
                "underlier": underlier, "product": str(product), "terms": terms,
                "tenor": (None if tenor in ("", None) else tenor), "notional": notional,
            })

with open("/tmp/wb_instruments.json", "w") as f:
    json.dump(out, f)
print(f"extracted {len(out)} INSTRUMENT cells -> /tmp/wb_instruments.json")
